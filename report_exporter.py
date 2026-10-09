
import re
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def markdown_to_word(markdown_text):
    """
    将 Research Agent 生成的 Markdown 报告转换为 Word 文档。
    返回 BytesIO，可直接供 Streamlit 下载。
    """

    document = Document()

    # 页面设置
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # 正文字体
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)

    for line in markdown_text.splitlines():

        line = line.rstrip()

        if not line.strip():
            continue

        # 一级标题
        if line.startswith("# "):
            paragraph = document.add_heading(
                line[2:].strip(), level=0
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 二级标题
        elif line.startswith("## "):
            document.add_heading(
                line[3:].strip(), level=1
            )

        # 三级标题
        elif line.startswith("### "):
            document.add_heading(
                line[4:].strip(), level=2
            )

        # 无序列表
        elif re.match(r"^\s*[-*]\s+", line):
            text = re.sub(r"^\s*[-*]\s+", "", line)
            add_formatted_paragraph(
                document, text, style="List Bullet"
            )

        # 有序列表
        elif re.match(r"^\s*\d+[.)]\s+", line):
            text = re.sub(r"^\s*\d+[.)]\s+", "", line)
            add_formatted_paragraph(
                document, text, style="List Number"
            )

        # 普通段落
        else:
            add_formatted_paragraph(document, line)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output


def add_formatted_paragraph(document, text, style=None):
    """
    简单处理 Markdown 粗体和行内代码。
    保留正文中的引用编号，例如 [T1-E1]。
    """

    paragraph = document.add_paragraph(style=style)

    pattern = r"(\*\*.*?\*\*|`.*?`)"

    parts = re.split(pattern, text)

    for part in parts:

        if not part:
            continue

        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True

        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"

        else:
            paragraph.add_run(part)

    return paragraph